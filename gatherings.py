"""Gatherings: one shared place to plan anything the house hosts.

A kid's birthday party is the first case (one a year per kid, from family.KID_BIRTHDAYS),
but the shape is the same for a family visit or a small New Year's thing: a thin header
(what, when, where), a PLAN of things to decide or do grouped by workstream (venue, food,
cake, invitations...), and a GUEST LIST with RSVPs. Both parents work the same rows, so
the list is the mental load made visible instead of held by one person.

Three tables, all kept by the parents, nothing derived from a feed:
  gatherings        -- the header + the post-mortem (how it went) once the day has passed
  gathering_items   -- the plan, one row per decision/task, by workstream
  gathering_guests  -- who is invited and whether they answered

Each KIND seeds a starter plan (STARTERS) so a new gathering opens with the questions
already asked, not a blank page. "Start next year's" clones a finished gathering with
every status reset and the dates a year on -- that is how a post-mortem turns into next
year's plan instead of a memory.

The gathering's date is mirrored into `events` (source='gathering', source_ref=slug) so
it sits in This Week and in Ask like everything else the family has coming.

    python gatherings.py --self-test
"""
from __future__ import annotations

import io
import re
import sqlite3
import sys
from datetime import date, datetime, timedelta

import db
import family

KINDS = {
    "kid_birthday": "Kid's birthday party",
    "family_visit": "Family visit (house guests)",
    "party": "Party or get-together",
}
STATUSES = ["open", "decided", "done", "skipped"]
STATUS_LABELS = {"open": "Open", "decided": "Decided", "done": "Done", "skipped": "Skipped"}
RSVPS = ["not_asked", "invited", "yes", "no", "maybe"]
RSVP_LABELS = {"not_asked": "Not invited yet", "invited": "No reply", "yes": "Yes", "no": "No",
               "maybe": "Maybe"}


def bind_household() -> None:
    """(Re)derive the household-dependent constants from family. Runs at import; the
    self-test calls it again after family.use_example().

    OWNERS: "" (nobody yet) + the parents. KID_BIRTHDAYS: month-day of each kid's birthday.
    The index nudges when a birthday is inside the planning window and no party exists for
    it -- the "one a year per kid, we always have to plan it" fact, kept in one place."""
    g = globals()
    g["OWNERS"] = [""] + list(family.PARENTS)
    g["KIDS"] = list(family.KIDS)
    g["KID_BIRTHDAYS"] = dict(family.KID_BIRTHDAYS)


bind_household()
NUDGE_DAYS = 120
HOME_WINDOW_DAYS = 90  # a venue is booked 6-8 weeks out; 30 (the trip's) is too late to help

# The starter plans. Every row is a question the family has to answer anyway; the point
# of listing them is that nobody has to remember to ask. Order = the order the answers
# tend to come in.
STARTERS: dict[str, list[tuple[str, list[str]]]] = {
    "kid_birthday": [
        ("Basics", ["Pick the party date and time", "Set the budget",
                    "Decide the size: how many kids, drop-off or parents stay"]),
        ("Venue", ["Choose the venue (home or a place)", "Book it and pay the deposit",
                   "Confirm what the venue includes (tables, staff, food, time slot)"]),
        ("Guests & invitations", ["Build the guest list", "Decide where invitations come from (Evite, paper, class list)",
                                  "Send the invitations", "Chase RSVPs", "Give the venue the final headcount"]),
        ("Food & drink", ["Kids' food", "Grown-ups' food", "Drinks", "Allergies to plan around"]),
        ("Cake", ["Choose the cake (bakery or homemade, flavor)", "Order it", "Candles, plates, forks, the knife",
                  "Pick it up"]),
        ("Theme & decorations", ["Pick the theme", "Decorations and balloons", "Tablecloths, plates, cups"]),
        ("Activities", ["Entertainment (games, a character, a craft)", "Music", "Run of show, hour by hour"]),
        ("Favors & thank-yous", ["Party favors", "Thank-you notes afterward"]),
        ("Day of", ["What to bring (the checklist)", "Who picks up the cake", "Photos",
                    "Keep a list of gifts for the thank-yous"]),
    ],
    "family_visit": [
        ("Basics", ["Arrival and departure dates", "Who is coming", "Airport or station pickups"]),
        ("Sleeping", ["Where everyone sleeps (guest room, air mattress, kids double up)",
                      "Linens, towels, pillows"]),
        ("Meals", ["First-night dinner", "Grocery run before they arrive", "Restaurants to book",
                   "Anything they don't eat"]),
        ("Kids' week", ["School and activity schedule during the visit", "Sitter or cover needed"]),
        ("Things to do", ["Activities with the kids", "A grown-ups-only evening"]),
        ("House", ["Guest room and bathroom ready", "Anything to borrow or buy (crib, car seat)"]),
    ],
    "party": [
        ("Basics", ["Pick the date and time", "Set the budget", "How many people"]),
        ("Guests & invitations", ["Build the guest list", "Send the invitations", "Chase RSVPs"]),
        ("Food & drink", ["Food (cook, cater, order in)", "Drinks", "Dessert", "Allergies to plan around"]),
        ("Setup", ["Music or playlist", "Decorations", "Tables, chairs, glassware"]),
        ("Kids", ["Bedtime plan for the kids", "Sitter for the evening"]),
        ("Day of", ["Shopping list", "Cleanup plan"]),
    ],
}
WORKSTREAM_ORDER = {kind: [w for w, _ in rows] for kind, rows in STARTERS.items()}


# ---------------------------------------------------------------- helpers


def slugify(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (name or "").lower()).strip("-")
    return s or "gathering"


def _d(iso: str | None) -> date | None:
    try:
        return datetime.strptime((iso or "")[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def anchor_date(g: dict) -> str | None:
    """What to count down to: the party date once chosen, the occasion until then."""
    return g.get("event_date") or g.get("occasion_date")


def _money(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip().replace("$", "").replace(",", "")
    if not s:
        return None
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def _int(v, default=0) -> int:
    try:
        return max(0, int(str(v).strip() or default))
    except (TypeError, ValueError):
        return default


# ---------------------------------------------------------------- gatherings


G_FIELDS = ("name", "kind", "honoree", "occasion_date", "event_date", "end_date", "start_time",
            "end_time", "venue", "headcount_goal", "notes")


def create(con: sqlite3.Connection, data: dict, who: str | None, seed: bool = True) -> dict:
    """New gathering. The slug carries the year of its anchor date so next year's clone
    never collides. Seeds the starter plan for its kind unless told not to."""
    clean = {k: (data.get(k) or "").strip() for k in G_FIELDS}
    if not clean["name"]:
        raise ValueError("name is required")
    if clean["kind"] not in KINDS:
        raise ValueError("kind must be one of " + ", ".join(KINDS))
    if clean["honoree"] and clean["honoree"] not in KIDS:
        raise ValueError("honoree must be one of " + ", ".join(KIDS))
    for k in ("occasion_date", "event_date", "end_date"):
        if clean[k] and not _d(clean[k]):
            raise ValueError(f"{k} must be YYYY-MM-DD")
    if not clean["occasion_date"] and not clean["event_date"]:
        raise ValueError("a date is required: the occasion or the party date")
    if not clean["occasion_date"]:
        clean["occasion_date"] = clean["event_date"]
    if clean["end_date"] and clean["event_date"] and clean["end_date"] < clean["event_date"]:
        raise ValueError("the last day is before the first day")
    year = (clean["event_date"] or clean["occasion_date"])[:4]
    base = f"{slugify(clean['name'])}-{year}"
    slug, n = base, 2
    while con.execute("SELECT 1 FROM gatherings WHERE slug=?", (slug,)).fetchone():
        slug, n = f"{base}-{n}", n + 1
    now = db.now()
    con.execute(
        "INSERT INTO gatherings (slug, name, kind, honoree, occasion_date, event_date, end_date, "
        "start_time, end_time, venue, headcount_goal, budget, notes, status, cloned_from, "
        "created_at, updated_at, updated_by) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (slug, clean["name"][:80], clean["kind"], clean["honoree"] or None,
         clean["occasion_date"], clean["event_date"] or None, clean["end_date"] or None,
         clean["start_time"] or None, clean["end_time"] or None, clean["venue"][:200] or None,
         clean["headcount_goal"][:120] or None, _money(data.get("budget")), clean["notes"][:1000] or None,
         "planning", data.get("cloned_from"), now, now, who))
    if seed:
        for s_i, (ws, titles) in enumerate(STARTERS[clean["kind"]]):
            for t_i, title in enumerate(titles):
                con.execute(
                    "INSERT INTO gathering_items (gathering, workstream, title, status, sort, updated_at) "
                    "VALUES (?,?,?,?,?,?)", (slug, ws, title, "open", s_i * 100 + t_i, now))
    con.commit()
    mirror_event(con, slug)
    return get(con, slug)


def get(con: sqlite3.Connection, slug: str) -> dict | None:
    r = con.execute("SELECT * FROM gatherings WHERE slug=?", (slug,)).fetchone()
    return dict(r) if r else None


def update(con: sqlite3.Connection, slug: str, data: dict, who: str | None) -> dict:
    """Header edits from the page. Only the keys present in `data` change, so a form with
    one field can never blank the others."""
    g = get(con, slug)
    if not g:
        raise ValueError("no such gathering")
    sets, args = [], []
    for k in G_FIELDS + ("status", "went_well", "do_differently", "attended"):
        if k in data:
            v = (data.get(k) or "").strip() or None
            if k == "kind" and v not in KINDS:
                raise ValueError("bad kind")
            if k == "honoree" and v and v not in KIDS:
                raise ValueError("bad honoree")
            if k == "status" and v not in ("planning", "done", "cancelled"):
                raise ValueError("bad status")
            if k in ("occasion_date", "event_date", "end_date") and v and not _d(v):
                raise ValueError(f"{k} must be YYYY-MM-DD")
            if k == "name" and not v:
                raise ValueError("name is required")
            sets.append(f"{k}=?")
            args.append(v)
    for k in ("budget", "actual_cost"):
        if k in data:
            sets.append(f"{k}=?")
            args.append(_money(data.get(k)))
    if not sets:
        return g
    sets += ["updated_at=?", "updated_by=?"]
    args += [db.now(), who, slug]
    con.execute(f"UPDATE gatherings SET {', '.join(sets)} WHERE slug=?", args)
    con.commit()
    mirror_event(con, slug)
    return get(con, slug)


def delete(con: sqlite3.Connection, slug: str) -> int:
    con.execute("DELETE FROM gathering_items WHERE gathering=?", (slug,))
    con.execute("DELETE FROM gathering_guests WHERE gathering=?", (slug,))
    con.execute("DELETE FROM events WHERE source='gathering' AND source_ref=?", (slug,))
    n = con.execute("DELETE FROM gatherings WHERE slug=?", (slug,)).rowcount
    con.commit()
    return n


def is_past(g: dict, today: date) -> bool:
    d = _d(g.get("end_date") or anchor_date(g))
    return bool(d and d < today)


def summary(con: sqlite3.Connection, g: dict, today: date) -> dict:
    """The numbers a card shows. Counted from the rows, never stored."""
    slug = g["slug"]
    items = con.execute("SELECT status, cost, paid FROM gathering_items WHERE gathering=?", (slug,)).fetchall()
    live = [i for i in items if i["status"] != "skipped"]
    n_open = sum(1 for i in live if i["status"] == "open")
    n_decided = sum(1 for i in live if i["status"] == "decided")
    n_done = sum(1 for i in live if i["status"] == "done")
    planned = round(sum(i["cost"] or 0 for i in live), 2)
    paid = round(sum(i["cost"] or 0 for i in live if i["paid"]), 2)
    guests = con.execute("SELECT adults, kids, rsvp FROM gathering_guests WHERE gathering=?", (slug,)).fetchall()
    invited = [x for x in guests if x["rsvp"] != "not_asked"]
    yes = [x for x in invited if x["rsvp"] == "yes"]
    d = _d(anchor_date(g))
    out = {
        "n_items": len(live), "n_open": n_open, "n_decided": n_decided, "n_done": n_done,
        "n_skipped": len(items) - len(live),
        "pct": int(round(100 * n_done / len(live))) if live else 0,
        "planned": planned, "paid": paid,
        "n_guests": len(guests), "n_invited": len(invited), "n_yes": len(yes),
        "n_no_reply": sum(1 for x in invited if x["rsvp"] == "invited"),
        "n_no": sum(1 for x in invited if x["rsvp"] == "no"),
        "n_maybe": sum(1 for x in invited if x["rsvp"] == "maybe"),
        "yes_adults": sum(x["adults"] or 0 for x in yes),
        "yes_kids": sum(x["kids"] or 0 for x in yes),
        "invited_adults": sum(x["adults"] or 0 for x in invited),
        "invited_kids": sum(x["kids"] or 0 for x in invited),
        "days_until": (d - today).days if d else None,
        "date_set": bool(g.get("event_date")),
        "past": is_past(g, today),
    }
    return out


def listing(con: sqlite3.Connection, today: date) -> dict:
    """Index: coming up (soonest first), past (most recent first), and the birthday nudges."""
    rows = [dict(r) for r in con.execute("SELECT * FROM gatherings")]
    for g in rows:
        g.update(summary(con, g, today))
        g["anchor"] = anchor_date(g)
        g["kind_label"] = KINDS.get(g["kind"], g["kind"])
    upcoming = sorted([g for g in rows if not g["past"] and g["status"] != "cancelled"],
                      key=lambda g: g["anchor"] or "")
    past = sorted([g for g in rows if g["past"] or g["status"] == "cancelled"],
                  key=lambda g: g["anchor"] or "", reverse=True)
    return {"upcoming": upcoming, "past": past, "nudges": birthday_nudges(rows, today)}


def birthday_nudges(rows: list[dict], today: date) -> list[dict]:
    """A kid's birthday inside the window with no party being planned for it. A kid's
    party each year is the one thing this must never let slip quietly."""
    out = []
    for kid, md in KID_BIRTHDAYS.items():
        for yr in (today.year, today.year + 1):
            bd = _d(f"{yr}-{md}")
            if not bd or bd < today or (bd - today).days > NUDGE_DAYS:
                continue
            covered = any(g["kind"] == "kid_birthday" and g["honoree"] == kid
                          and (g.get("occasion_date") or "")[:4] == str(yr)
                          and g["status"] != "cancelled" for g in rows)
            if not covered:
                out.append({"kid": kid, "date": bd.isoformat(), "days": (bd - today).days,
                            "name": f"{kid}'s birthday party"})
            break
    return out


def next_for_home(con: sqlite3.Connection, today: date) -> dict | None:
    """The one gathering the home page mentions: the soonest inside the window."""
    lst = listing(con, today)
    for g in lst["upcoming"]:
        if g["days_until"] is not None and g["days_until"] <= HOME_WINDOW_DAYS:
            return g
    return None


# ---------------------------------------------------------------- the plan


def plan(con: sqlite3.Connection, slug: str, kind: str | None = None) -> list[dict]:
    """Items grouped by workstream in the kind's canonical order, parent-added workstreams
    after. Each group carries its own done count."""
    rows = [dict(r) for r in con.execute(
        "SELECT * FROM gathering_items WHERE gathering=? ORDER BY sort, id", (slug,))]
    order = WORKSTREAM_ORDER.get(kind or "", [])
    groups: dict[str, dict] = {}
    for r in rows:
        grp = groups.setdefault(r["workstream"], {"workstream": r["workstream"], "items": [],
                                                  "done": 0, "live": 0})
        grp["items"].append(r)
        if r["status"] != "skipped":
            grp["live"] += 1
            if r["status"] == "done":
                grp["done"] += 1

    def key(ws):
        return (order.index(ws) if ws in order else len(order), ws.lower())
    return [groups[w] for w in sorted(groups, key=key)]


def add_item(con: sqlite3.Connection, slug: str, workstream: str, title: str, who: str | None,
             owner: str | None = None, due: str | None = None) -> dict:
    workstream = (workstream or "Other").strip()[:60]
    title = (title or "").strip()[:200]
    if not title:
        raise ValueError("title is required")
    g = get(con, slug)
    if not g:
        raise ValueError("no such gathering")
    order = WORKSTREAM_ORDER.get(g["kind"], [])
    base = (order.index(workstream) if workstream in order else len(order)) * 100
    top = con.execute("SELECT COALESCE(MAX(sort), ?) FROM gathering_items WHERE gathering=? AND workstream=?",
                      (base - 1, slug, workstream)).fetchone()[0]
    cur = con.execute(
        "INSERT INTO gathering_items (gathering, workstream, title, status, owner, due, sort, updated_at, updated_by) "
        "VALUES (?,?,?,?,?,?,?,?,?)",
        (slug, workstream, title, "open", owner or None, (due or "").strip()[:10] or None, top + 1, db.now(), who))
    con.commit()
    return dict(con.execute("SELECT * FROM gathering_items WHERE id=?", (cur.lastrowid,)).fetchone())


ITEM_FIELDS = ("title", "status", "owner", "due", "decision", "cost", "paid", "workstream")


def set_item(con: sqlite3.Connection, item_id: int, data: dict, who: str | None) -> dict:
    """Edit any subset of an item's fields. Setting a decision on an OPEN item moves it to
    decided -- writing the answer down is the decision -- unless the caller set status too."""
    r = con.execute("SELECT * FROM gathering_items WHERE id=?", (item_id,)).fetchone()
    if not r:
        raise ValueError("no such item")
    sets, args = [], []
    if "status" in data:
        st = (data.get("status") or "").strip()
        if st not in STATUSES:
            raise ValueError("status must be one of " + ", ".join(STATUSES))
        sets.append("status=?"); args.append(st)
    if "title" in data:
        t = (data.get("title") or "").strip()[:200]
        if not t:
            raise ValueError("title is required")
        sets.append("title=?"); args.append(t)
    if "workstream" in data:
        sets.append("workstream=?"); args.append((data.get("workstream") or "Other").strip()[:60])
    if "owner" in data:
        o = (data.get("owner") or "").strip()
        if o and o not in OWNERS[1:]:
            raise ValueError("owner must be " + " or ".join(OWNERS[1:]))
        sets.append("owner=?"); args.append(o or None)
    if "due" in data:
        d = (data.get("due") or "").strip()[:10]
        if d and not _d(d):
            raise ValueError("due must be YYYY-MM-DD")
        sets.append("due=?"); args.append(d or None)
    if "decision" in data:
        dec = (data.get("decision") or "").strip()[:600]
        sets.append("decision=?"); args.append(dec or None)
        if dec and "status" not in data and r["status"] == "open":
            sets.append("status=?"); args.append("decided")
    if "cost" in data:
        sets.append("cost=?"); args.append(_money(data.get("cost")))
    if "paid" in data:
        sets.append("paid=?"); args.append(1 if str(data.get("paid")) in ("1", "true", "True", "on", "yes") else 0)
    if not sets:
        return dict(r)
    sets += ["updated_at=?", "updated_by=?"]
    args += [db.now(), who, item_id]
    con.execute(f"UPDATE gathering_items SET {', '.join(sets)} WHERE id=?", args)
    con.commit()
    return dict(con.execute("SELECT * FROM gathering_items WHERE id=?", (item_id,)).fetchone())


def delete_item(con: sqlite3.Connection, item_id: int) -> int:
    n = con.execute("DELETE FROM gathering_items WHERE id=?", (item_id,)).rowcount
    con.commit()
    return n


# ---------------------------------------------------------------- guests


def guests(con: sqlite3.Connection, slug: str) -> list[dict]:
    return [dict(r) for r in con.execute(
        "SELECT * FROM gathering_guests WHERE gathering=? ORDER BY sort, id", (slug,))]


def add_guest(con: sqlite3.Connection, slug: str, data: dict, who: str | None) -> dict:
    name = (data.get("name") or "").strip()[:120]
    if not name:
        raise ValueError("name is required")
    rsvp = (data.get("rsvp") or "not_asked").strip()
    if rsvp not in RSVPS:
        raise ValueError("rsvp must be one of " + ", ".join(RSVPS))
    top = con.execute("SELECT COALESCE(MAX(sort), 0) FROM gathering_guests WHERE gathering=?", (slug,)).fetchone()[0]
    cur = con.execute(
        "INSERT INTO gathering_guests (gathering, name, household, adults, kids, rsvp, contact, notes, sort, "
        "updated_at, updated_by) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (slug, name, (data.get("household") or "").strip()[:120] or None, _int(data.get("adults"), 0),
         _int(data.get("kids"), 1), rsvp, (data.get("contact") or "").strip()[:200] or None,
         (data.get("notes") or "").strip()[:400] or None, top + 1, db.now(), who))
    con.commit()
    return dict(con.execute("SELECT * FROM gathering_guests WHERE id=?", (cur.lastrowid,)).fetchone())


def add_guests_from_text(con: sqlite3.Connection, slug: str, text: str, who: str | None,
                         adults: int = 0, kids: int = 1) -> int:
    """Paste the class list. One guest per line; a line like "Liam - Sarah 555-0100" keeps
    the part after the dash as the contact. Duplicate names (case-insensitive) are skipped
    so pasting the same list twice adds nothing."""
    have = {g["name"].strip().lower() for g in guests(con, slug)}
    n = 0
    for line in (text or "").splitlines():
        line = line.strip().lstrip("-*• ").strip()
        if not line:
            continue
        name, contact = line, ""
        m = re.match(r"^(.*?)\s+[-–—:]\s+(.*)$", line)
        if m:
            name, contact = m.group(1).strip(), m.group(2).strip()
        if name.lower() in have:
            continue
        add_guest(con, slug, {"name": name, "contact": contact, "adults": adults, "kids": kids}, who)
        have.add(name.lower())
        n += 1
    return n


GUEST_FIELDS = ("name", "household", "adults", "kids", "rsvp", "contact", "notes")


def set_guest(con: sqlite3.Connection, guest_id: int, data: dict, who: str | None) -> dict:
    r = con.execute("SELECT * FROM gathering_guests WHERE id=?", (guest_id,)).fetchone()
    if not r:
        raise ValueError("no such guest")
    sets, args = [], []
    for k in GUEST_FIELDS:
        if k not in data:
            continue
        v = data.get(k)
        if k == "name":
            v = (v or "").strip()[:120]
            if not v:
                raise ValueError("name is required")
        elif k in ("adults", "kids"):
            v = _int(v, 0)
        elif k == "rsvp":
            v = (v or "not_asked").strip()
            if v not in RSVPS:
                raise ValueError("bad rsvp")
        else:
            v = (v or "").strip()[:400] or None
        sets.append(f"{k}=?"); args.append(v)
    if not sets:
        return dict(r)
    sets += ["updated_at=?", "updated_by=?"]
    args += [db.now(), who, guest_id]
    con.execute(f"UPDATE gathering_guests SET {', '.join(sets)} WHERE id=?", args)
    con.commit()
    return dict(con.execute("SELECT * FROM gathering_guests WHERE id=?", (guest_id,)).fetchone())


def delete_guest(con: sqlite3.Connection, guest_id: int) -> int:
    n = con.execute("DELETE FROM gathering_guests WHERE id=?", (guest_id,)).rowcount
    con.commit()
    return n


def mark_all_invited(con: sqlite3.Connection, slug: str, who: str | None) -> int:
    """'Invitations went out' in one press: every not-yet-asked guest becomes invited/no reply."""
    n = con.execute("UPDATE gathering_guests SET rsvp='invited', updated_at=?, updated_by=? "
                    "WHERE gathering=? AND rsvp='not_asked'", (db.now(), who, slug)).rowcount
    con.commit()
    return n


# ---------------------------------------------------------------- clone + mirror


def clone(con: sqlite3.Connection, slug: str, who: str | None, years: int = 1) -> dict:
    """Next year's, started from this one: the same plan with every status reset and the
    decisions kept as last year's answers in the notes; the same guest list with RSVPs
    cleared. Dates move `years` on; the party date is left for them to pick again."""
    g = get(con, slug)
    if not g:
        raise ValueError("no such gathering")

    def shift(iso):
        d = _d(iso)
        if not d:
            return None
        try:
            return d.replace(year=d.year + years).isoformat()
        except ValueError:  # Feb 29
            return d.replace(year=d.year + years, day=28).isoformat()
    name = re.sub(r"\b(19|20)\d{2}\b", "", g["name"]).strip() or g["name"]
    new = create(con, {
        "name": name, "kind": g["kind"], "honoree": g["honoree"],
        "occasion_date": shift(g["occasion_date"]), "event_date": "",
        "venue": g["venue"], "headcount_goal": g["headcount_goal"], "budget": g["budget"],
        "notes": g["notes"], "cloned_from": slug,
    }, who, seed=False)
    now = db.now()
    for it in con.execute("SELECT * FROM gathering_items WHERE gathering=? ORDER BY sort, id", (slug,)):
        last = (f"Last time: {it['decision']}" if it["decision"] else None)
        con.execute(
            "INSERT INTO gathering_items (gathering, workstream, title, status, owner, decision, cost, paid, sort, updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?)",
            (new["slug"], it["workstream"], it["title"], "skipped" if it["status"] == "skipped" else "open",
             it["owner"], last, it["cost"], 0, it["sort"], now))
    for gu in con.execute("SELECT * FROM gathering_guests WHERE gathering=? ORDER BY sort, id", (slug,)):
        con.execute(
            "INSERT INTO gathering_guests (gathering, name, household, adults, kids, rsvp, contact, notes, sort, updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?)",
            (new["slug"], gu["name"], gu["household"], gu["adults"], gu["kids"], "not_asked", gu["contact"],
             gu["notes"], gu["sort"], now))
    con.commit()
    return get(con, new["slug"])


def mirror_event(con: sqlite3.Connection, slug: str) -> None:
    """One `events` row per gathering with a party date, so it shows in This Week and Ask.
    Identity is (source='gathering', source_ref=slug); the title may change freely. No
    party date yet = no row (the occasion itself is on the calendar already as a birthday)."""
    g = get(con, slug)
    if not g:
        return
    rows = con.execute("SELECT id FROM events WHERE source='gathering' AND source_ref=?", (slug,)).fetchall()
    if not g.get("event_date") or g.get("status") == "cancelled":
        for r in rows:
            con.execute("DELETE FROM events WHERE id=?", (r["id"],))
        con.commit()
        return
    kid = g["honoree"] or "Both"
    title = g["name"] + (f" at {g['venue']}" if g.get("venue") else "")
    details = f"Planned on Family Manager: /gathering/{slug}"
    now = db.now()
    if rows:
        con.execute("UPDATE events SET kid=?, title=?, event_date=?, end_date=?, start_time=?, end_time=?, "
                    "details=?, status='active', last_seen_at=? WHERE id=?",
                    (kid, title, g["event_date"], g.get("end_date"), g.get("start_time"), g.get("end_time"),
                     details, now, rows[0]["id"]))
        for r in rows[1:]:
            con.execute("DELETE FROM events WHERE id=?", (r["id"],))
    else:
        # UNIQUE(kid, title, event_date) may already hold a copy from a calendar feed;
        # INSERT OR IGNORE keeps that copy and this gathering simply has no mirror row.
        con.execute("INSERT OR IGNORE INTO events (kid, type, title, event_date, end_date, start_time, end_time, "
                    "category, details, source, source_ref, status, created_at, last_seen_at) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (kid, "event", title, g["event_date"], g.get("end_date"), g.get("start_time"), g.get("end_time"),
                     "travel_home", details, "gathering", slug, "active", now, now))
    con.commit()


# ---------------------------------------------------------------- ask + workbook


def ask_lines(con: sqlite3.Connection, today: date) -> list[str]:
    """What Ask reads: every gathering as one line, with its open items and RSVP tally."""
    out = []
    for g in listing(con, today)["upcoming"] + listing(con, today)["past"][:3]:
        s = f"{g['name']} ({g['kind_label']})"
        if g["event_date"]:
            s += f", on {g['event_date']}" + (f" {g['start_time']}" if g["start_time"] else "")
        else:
            s += f", date not picked yet (occasion {g['occasion_date']})"
        if g["venue"]:
            s += f", at {g['venue']}"
        s += f"; plan {g['n_done']} done / {g['n_decided']} decided / {g['n_open']} open of {g['n_items']}"
        if g["n_invited"]:
            s += (f"; guests: {g['n_invited']} invited, {g['n_yes']} yes ({g['yes_adults']} adults + "
                  f"{g['yes_kids']} kids), {g['n_no_reply']} no reply, {g['n_no']} no")
        elif g["n_guests"]:
            s += f"; guest list has {g['n_guests']} names, invitations not sent"
        if g["budget"]:
            s += f"; budget ${g['budget']:,.0f}, planned ${g['planned']:,.0f}"
        out.append(s)
        opens = con.execute("SELECT workstream, title, owner, due FROM gathering_items WHERE gathering=? AND status='open' "
                            "ORDER BY sort LIMIT 12", (g["slug"],)).fetchall()
        for o in opens:
            out.append(f"  - open: {o['workstream']}: {o['title']}" + (f" ({o['owner']})" if o["owner"] else "")
                       + (f", due {o['due']}" if o["due"] else ""))
        decided = con.execute("SELECT workstream, title, decision FROM gathering_items WHERE gathering=? "
                              "AND decision IS NOT NULL ORDER BY sort LIMIT 12", (g["slug"],)).fetchall()
        for o in decided:
            out.append(f"  - {o['workstream']}: {o['title']} -> {o['decision']}")
    return out


def workbook_bytes(con: sqlite3.Connection, slug: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sheet(ws, cols, rows, widths):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C4974C")
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    g = get(con, slug) or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    sheet(ws, ["Workstream", "Item", "Status", "Owner", "Due", "Decision", "Cost", "Paid", "Updated by"],
          [[r["workstream"], r["title"], STATUS_LABELS.get(r["status"], r["status"]), r["owner"] or "",
            r["due"] or "", r["decision"] or "", r["cost"], "yes" if r["paid"] else "", r["updated_by"] or ""]
           for grp in plan(con, slug, g.get("kind")) for r in grp["items"]],
          [22, 44, 10, 8, 12, 50, 10, 6, 10])
    ws2 = wb.create_sheet("Guests")
    sheet(ws2, ["Name", "Household", "Adults", "Kids", "RSVP", "Contact", "Notes"],
          [[r["name"], r["household"] or "", r["adults"], r["kids"], RSVP_LABELS.get(r["rsvp"], r["rsvp"]),
            r["contact"] or "", r["notes"] or ""] for r in guests(con, slug)],
          [28, 20, 8, 8, 14, 26, 36])
    ws3 = wb.create_sheet("Gathering")
    facts = [("Name", g.get("name")), ("Kind", KINDS.get(g.get("kind"), g.get("kind"))), ("For", g.get("honoree")),
             ("Occasion", g.get("occasion_date")), ("Party date", g.get("event_date")), ("Ends", g.get("end_date")),
             ("Time", " - ".join(x for x in (g.get("start_time"), g.get("end_time")) if x)),
             ("Venue", g.get("venue")), ("Headcount goal", g.get("headcount_goal")), ("Budget", g.get("budget")),
             ("Notes", g.get("notes")), ("Status", g.get("status")), ("What went well", g.get("went_well")),
             ("Do differently", g.get("do_differently")), ("Actual cost", g.get("actual_cost")),
             ("Who came", g.get("attended"))]
    sheet(ws3, ["Field", "Value"], [[k, v if v is not None else ""] for k, v in facts], [18, 70])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def index_workbook_bytes(con: sqlite3.Connection, today: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Gatherings"
    cols = ["Name", "Kind", "For", "Occasion", "Party date", "Venue", "Status", "Open", "Decided", "Done",
            "Invited", "Yes", "No reply", "Budget", "Planned", "Actual"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="C4974C")
    lst = listing(con, today)
    for g in lst["upcoming"] + lst["past"]:
        ws.append([g["name"], g["kind_label"], g["honoree"] or "", g["occasion_date"], g["event_date"] or "",
                   g["venue"] or "", g["status"], g["n_open"], g["n_decided"], g["n_done"], g["n_invited"],
                   g["n_yes"], g["n_no_reply"], g["budget"], g["planned"], g["actual_cost"]])
    for i, w in enumerate([30, 22, 8, 12, 12, 24, 10, 7, 8, 7, 8, 6, 9, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- self-test


def _self_test() -> int:
    fails = 0

    def check(name, ok):
        nonlocal fails
        print(("  PASS  " if ok else "  FAIL  ") + name)
        if not ok:
            fails += 1

    # The example family: Sam + Jordan, Ava (birthday 03-14) and Leo (birthday 10-02).
    family.use_example()
    bind_household()

    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript(db.SCHEMA)
    # 60 days before Leo's birthday, the same distance the original fixture used.
    today = date(2026, 8, 3)

    try:
        create(con, {"name": "x", "kind": "kid_birthday"}, "Sam")
        check("a gathering needs a date", False)
    except ValueError:
        check("a gathering needs a date", True)
    try:
        create(con, {"name": "x", "kind": "picnic", "occasion_date": "2026-10-02"}, "Sam")
        check("kind is validated", False)
    except ValueError:
        check("kind is validated", True)

    g = create(con, {"name": "Leo's birthday party", "kind": "kid_birthday", "honoree": "Leo",
                     "occasion_date": "2026-10-02", "budget": "$400"}, "Sam")
    check("slug carries the year", g["slug"] == "leo-s-birthday-party-2026")
    check("budget parses money text", g["budget"] == 400.0)
    n_seed = sum(len(t) for _, t in STARTERS["kid_birthday"])
    check("starter plan seeded for the kind",
          con.execute("SELECT COUNT(*) FROM gathering_items WHERE gathering=?", (g["slug"],)).fetchone()[0] == n_seed)
    p = plan(con, g["slug"], g["kind"])
    check("plan groups in the starter's order", [x["workstream"] for x in p] == WORKSTREAM_ORDER["kid_birthday"])
    check("no party date = no events mirror row",
          con.execute("SELECT COUNT(*) FROM events WHERE source='gathering'").fetchone()[0] == 0)
    s = summary(con, g, today)
    check("summary counts down to the occasion while the date is unset",
          s["days_until"] == 60 and not s["date_set"] and s["n_open"] == n_seed and s["pct"] == 0)

    g2 = create(con, {"name": "Leo's birthday party", "kind": "kid_birthday", "honoree": "Leo",
                      "occasion_date": "2026-10-02"}, "Sam")
    check("same name same year gets a distinct slug", g2["slug"] == "leo-s-birthday-party-2026-2")
    delete(con, g2["slug"])
    check("delete removes the plan too",
          con.execute("SELECT COUNT(*) FROM gathering_items WHERE gathering=?", (g2["slug"],)).fetchone()[0] == 0)

    # The plan.
    first = p[0]["items"][0]
    it = set_item(con, first["id"], {"decision": "Sat Oct 3, 2-4pm"}, "Jordan")
    check("writing a decision on an open item makes it decided", it["status"] == "decided" and it["updated_by"] == "Jordan")
    it = set_item(con, first["id"], {"status": "done"}, "Jordan")
    check("status can be set outright", it["status"] == "done")
    it = set_item(con, first["id"], {"decision": ""}, "Jordan")
    check("clearing a decision does not move a done item back", it["status"] == "done" and it["decision"] is None)
    try:
        set_item(con, first["id"], {"status": "later"}, "Sam")
        check("bad status refused", False)
    except ValueError:
        check("bad status refused", True)
    try:
        set_item(con, first["id"], {"owner": "Grandma"}, "Sam")
        check("owner limited to the parents", False)
    except ValueError:
        check("owner limited to the parents", True)
    venue = next(x for x in p if x["workstream"] == "Venue")["items"]
    set_item(con, venue[0]["id"], {"decision": "Bounce U", "cost": "350", "paid": "1"}, "Sam")
    set_item(con, venue[1]["id"], {"status": "skipped"}, "Sam")
    set_item(con, venue[2]["id"], {"cost": "$25.50"}, "Sam")
    s = summary(con, g, today)
    check("planned and paid sum only live rows", s["planned"] == 375.5 and s["paid"] == 350.0 and s["n_skipped"] == 1)
    extra = add_item(con, g["slug"], "Venue", "Ask about the parking", "Jordan", owner="Jordan", due="2026-09-01")
    vgrp = next(x for x in plan(con, g["slug"], g["kind"]) if x["workstream"] == "Venue")
    check("a parent-added item lands last in its workstream", vgrp["items"][-1]["id"] == extra["id"])
    new_ws = add_item(con, g["slug"], "Pinata", "Buy one", "Sam")
    check("a new workstream sorts after the starter ones",
          plan(con, g["slug"], g["kind"])[-1]["workstream"] == "Pinata")
    try:
        add_item(con, g["slug"], "Venue", "  ", "Sam")
        check("blank item refused", False)
    except ValueError:
        check("blank item refused", True)
    delete_item(con, new_ws["id"])

    # Guests.
    a = add_guest(con, g["slug"], {"name": "Liam", "household": "The Parks", "adults": 2, "kids": 1}, "Jordan")
    n = add_guests_from_text(con, g["slug"], "Liam\n- Noah - Jess 555-0100\n\nOlivia\n", "Jordan")
    check("paste skips a duplicate and keeps the contact after the dash", n == 2 and
          next(x for x in guests(con, g["slug"]) if x["name"] == "Noah")["contact"] == "Jess 555-0100")
    check("paste twice adds nothing", add_guests_from_text(con, g["slug"], "Liam\nNoah\nOlivia", "Jordan") == 0)
    s = summary(con, g, today)
    check("names on the list are not yet invited", s["n_guests"] == 3 and s["n_invited"] == 0)
    check("mark all invited flips the not-asked rows", mark_all_invited(con, g["slug"], "Sam") == 3)
    set_guest(con, a["id"], {"rsvp": "yes"}, "Jordan")
    noah = next(x for x in guests(con, g["slug"]) if x["name"] == "Noah")
    set_guest(con, noah["id"], {"rsvp": "no"}, "Jordan")
    s = summary(con, g, today)
    check("rsvp tally: yes counts adults + kids, no reply is what is left",
          s["n_yes"] == 1 and s["yes_adults"] == 2 and s["yes_kids"] == 1 and s["n_no"] == 1 and s["n_no_reply"] == 1)
    try:
        set_guest(con, a["id"], {"rsvp": "perhaps"}, "Jordan")
        check("bad rsvp refused", False)
    except ValueError:
        check("bad rsvp refused", True)

    # Header edits + the mirror.
    g = update(con, g["slug"], {"event_date": "2026-10-03", "start_time": "14:00", "venue": "Bounce U"}, "Jordan")
    ev = con.execute("SELECT * FROM events WHERE source='gathering' AND source_ref=?", (g["slug"],)).fetchall()
    check("a party date creates ONE events row for the honoree",
          len(ev) == 1 and ev[0]["kid"] == "Leo" and ev[0]["event_date"] == "2026-10-03"
          and ev[0]["start_time"] == "14:00" and "Bounce U" in ev[0]["title"])
    g = update(con, g["slug"], {"event_date": "2026-10-04"}, "Jordan")
    ev = con.execute("SELECT event_date FROM events WHERE source='gathering' AND source_ref=?", (g["slug"],)).fetchall()
    check("moving the date updates the same row", len(ev) == 1 and ev[0][0] == "2026-10-04")
    check("a single-field update leaves the rest alone", g["venue"] == "Bounce U" and g["budget"] == 400.0)
    g = update(con, g["slug"], {"event_date": ""}, "Jordan")
    check("clearing the date removes the mirror row",
          con.execute("SELECT COUNT(*) FROM events WHERE source='gathering'").fetchone()[0] == 0)
    update(con, g["slug"], {"event_date": "2026-10-04"}, "Jordan")
    try:
        update(con, g["slug"], {"name": ""}, "Sam")
        check("name cannot be blanked", False)
    except ValueError:
        check("name cannot be blanked", True)

    # Index, nudges, home.
    lst = listing(con, today)
    check("index: upcoming has Leo, past is empty", [x["slug"] for x in lst["upcoming"]] == [g["slug"]] and not lst["past"])
    check("Leo's party covers his birthday; Ava's March birthday is outside the window",
          lst["nudges"] == [])
    check("home page: 95 days out is outside the 90-day window", next_for_home(con, date(2026, 7, 1)) is None)
    check("home page picks it inside the window", next_for_home(con, today)["slug"] == g["slug"])
    delete(con, g["slug"])
    lst = listing(con, today)
    check("with no party, Leo's birthday nudges",
          [(n["kid"], n["date"], n["days"]) for n in lst["nudges"]] == [("Leo", "2026-10-02", 60)])
    check("Ava's birthday nudges once it is inside the window",
          [n["kid"] for n in listing(con, date(2027, 1, 1))["nudges"]] == ["Ava"])

    # A past gathering, then next year's from it.
    m = create(con, {"name": "Ava's birthday party", "kind": "kid_birthday", "honoree": "Ava",
                     "occasion_date": "2026-03-14", "event_date": "2026-03-14"}, "Sam")
    cake = next(x for x in plan(con, m["slug"], m["kind"]) if x["workstream"] == "Cake")["items"]
    set_item(con, cake[0]["id"], {"decision": "Costco sheet cake, vanilla", "status": "done", "cost": "30"}, "Jordan")
    set_item(con, cake[1]["id"], {"status": "skipped"}, "Jordan")
    add_guest(con, m["slug"], {"name": "Mia", "adults": 1, "kids": 1, "rsvp": "yes"}, "Jordan")
    update(con, m["slug"], {"status": "done", "went_well": "Pool was a hit", "actual_cost": "512.40"}, "Sam")
    lst = listing(con, today)
    check("a finished gathering lists under past", [x["slug"] for x in lst["past"]] == [m["slug"]])
    nxt = clone(con, m["slug"], "Sam")
    check("clone: next year's slug, dates a year on, party date left open",
          nxt["slug"] == "ava-s-birthday-party-2027" and nxt["occasion_date"] == "2027-03-14"
          and nxt["event_date"] is None and nxt["cloned_from"] == m["slug"])
    ncake = next(x for x in plan(con, nxt["slug"], nxt["kind"]) if x["workstream"] == "Cake")["items"]
    check("clone: statuses reset, last year's answer kept as a note, skipped stays skipped",
          ncake[0]["status"] == "open" and ncake[0]["decision"] == "Last time: Costco sheet cake, vanilla"
          and ncake[0]["paid"] == 0 and ncake[1]["status"] == "skipped")
    check("clone: guests carried with RSVPs cleared",
          [(x["name"], x["rsvp"]) for x in guests(con, nxt["slug"])] == [("Mia", "not_asked")])
    check("clone: no nudge for Ava 2027 now", [n["kid"] for n in listing(con, date(2027, 1, 1))["nudges"]] == [])
    check("clone: the original is untouched", get(con, m["slug"])["status"] == "done"
          and con.execute("SELECT status FROM gathering_items WHERE id=?", (cake[0]["id"],)).fetchone()[0] == "done")

    lines = ask_lines(con, today)
    check("ask lines name the gathering, the plan tally and the answer",
          any("Ava's birthday party" in l and "done" in l for l in lines)
          and any("Costco sheet cake" in l for l in lines))

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(workbook_bytes(con, m["slug"])))
    check("workbook: Plan + Guests + Gathering sheets with rows",
          wb.sheetnames == ["Plan", "Guests", "Gathering"] and wb["Plan"].max_row == n_seed + 1 and wb["Guests"].max_row == 2)
    wb2 = load_workbook(io.BytesIO(index_workbook_bytes(con, today)))
    check("index workbook lists every gathering", wb2.active.max_row == 3)

    print("\n" + ("ALL PASS" if not fails else f"{fails} FAILED"))
    return 1 if fails else 0


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        sys.exit(_self_test())
    print(__doc__)
