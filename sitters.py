"""Babysitters in rotation + the sheet a sitter needs (pediatrician, allergies, the door code).

Two tables, both kept by the parents, nothing derived from a feed:
  sitters      -- who we call, in what order, and how to reach them
  sitter_info  -- one row per fact a sitter needs, grouped by section. Values are EMPTY
                  until a parent fills them in. A plausible default in a pediatrician
                  field is worse than a blank one, so the seed writes labels only.

Days off ties in: coverage_plan.coverage = 'Babysitter' + coverage_plan.sitter_id names
which sitter has that day, and the calendar event carries their name and number.

    python sitters.py --self-test
"""
from __future__ import annotations

import io
import sqlite3
import sys

import db
import family

BABYSITTER = "Babysitter"

# The sections a kid gets on the sheet, worded so they fit any kid.
KID_LABELS = ["Allergies", "Medications", "Bedtime + routine", "Food they eat / won't eat",
              "What calms them down", "Diapers / potty, if any"]


def build_sections() -> list[tuple[str, list[str]]]:
    """The info sheet, section by section, in the order a sitter reads it. "Reach us" has
    one cell line per parent, then one section per kid. Labels only -- the one pre-filled
    value is Poison Control, which is a national constant, not a guess."""
    return ([("Reach us", [f"{p} cell" for p in family.PARENTS]
              + ["Home address", "Where we are tonight"]),
             ("Pediatrician", ["Practice", "Phone", "After-hours line", "Address"]),
             ("Dentist", ["Practice", "Phone", "Address"]),
             ("Emergency", ["Poison Control", "Nearest ER", "Emergency contact 1 (name + phone)",
                            "Emergency contact 2 (name + phone)",
                            "Health insurance (carrier + member id)"])]
            + [(kid, list(KID_LABELS)) for kid in family.KIDS]
            + [("House", ["Getting in (door code / key)", "Alarm", "Wifi",
                          "First-aid kit + flashlight", "Thermostat", "Pets"]),
               ("Rules", ["Screen time", "Snacks", "Who may pick up",
                          "Who may NOT come to the door", "Bath night", "Anything else"])])


def bind_household() -> None:
    """(Re)derive KIDS_OK and SECTIONS from family. Runs at import; the self-test calls it
    again after family.use_example(). "Both" is the stored value for every kid."""
    g = globals()
    g["KIDS_OK"] = ["Both"] + list(family.KIDS)
    g["SECTIONS"] = build_sections()


bind_household()
PREFILLED = {("Emergency", "Poison Control"): "1-800-222-1222 (national, 24/7)"}


def seed_info(con: sqlite3.Connection) -> int:
    """Idempotent: adds any missing label, never touches a value a parent has entered."""
    n = 0
    for s_i, (section, labels) in enumerate(SECTIONS):
        for l_i, label in enumerate(labels):
            cur = con.execute(
                "INSERT OR IGNORE INTO sitter_info (section, label, value, sort) VALUES (?,?,?,?)",
                (section, label, PREFILLED.get((section, label)), s_i * 100 + l_i))
            n += cur.rowcount
    con.commit()
    return n


def info_sections(con: sqlite3.Connection) -> list[dict]:
    rows = con.execute("SELECT * FROM sitter_info ORDER BY sort, id").fetchall()
    out: list[dict] = []
    for r in rows:
        if not out or out[-1]["section"] != r["section"]:
            out.append({"section": r["section"], "rows": [], "filled": 0})
        out[-1]["rows"].append(dict(r))
        if (r["value"] or "").strip():
            out[-1]["filled"] += 1
    return out


def set_info(con: sqlite3.Connection, info_id: int, value: str, who: str | None) -> dict | None:
    con.execute("UPDATE sitter_info SET value=?, updated_at=?, updated_by=? WHERE id=?",
                ((value or "").strip()[:600], db.now(), who, info_id))
    con.commit()
    r = con.execute("SELECT * FROM sitter_info WHERE id=?", (info_id,)).fetchone()
    return dict(r) if r else None


def add_info_row(con: sqlite3.Connection, section: str, label: str, who: str | None) -> dict:
    """A parent can add a fact the seed never thought of; it lands last in its section."""
    section = (section or "Anything else").strip()[:60]
    label = (label or "").strip()[:80]
    if not label:
        raise ValueError("label is required")
    top = con.execute("SELECT COALESCE(MAX(sort), 0) FROM sitter_info WHERE section=?",
                      (section,)).fetchone()[0]
    base = next((i * 100 for i, (s, _) in enumerate(SECTIONS) if s == section), len(SECTIONS) * 100)
    con.execute("INSERT OR IGNORE INTO sitter_info (section, label, sort, updated_at, updated_by) "
                "VALUES (?,?,?,?,?)", (section, label, max(top, base) + 1, db.now(), who))
    con.commit()
    return dict(con.execute("SELECT * FROM sitter_info WHERE section=? AND label=?",
                            (section, label)).fetchone())


def roster(con: sqlite3.Connection, include_inactive: bool = True) -> list[dict]:
    q = "SELECT * FROM sitters"
    if not include_inactive:
        q += " WHERE active=1"
    q += " ORDER BY active DESC, sort, name COLLATE NOCASE"
    return [dict(r) for r in con.execute(q)]


SITTER_FIELDS = ("name", "phone", "email", "rate", "how_we_know", "kids_ok",
                 "availability", "notes")


def save_sitter(con: sqlite3.Connection, data: dict, who: str | None,
                sitter_id: int | None = None) -> dict:
    """Create or update one sitter. Refuses a blank name -- an unnamed row on a roster
    is a row nobody can call."""
    clean = {k: (data.get(k) or "").strip()[:300] for k in SITTER_FIELDS}
    if not clean["name"]:
        raise ValueError("name is required")
    if clean["kids_ok"] and clean["kids_ok"] not in KIDS_OK:
        raise ValueError("kids_ok must be one of " + ", ".join(KIDS_OK))
    active = 0 if str(data.get("active", "1")) in ("0", "false", "False", "") else 1
    now = db.now()
    if sitter_id is None:
        top = con.execute("SELECT COALESCE(MAX(sort), 0) FROM sitters").fetchone()[0]
        cur = con.execute(
            "INSERT INTO sitters (name, phone, email, rate, how_we_know, kids_ok, availability, "
            "notes, active, sort, updated_at, updated_by) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (*[clean[k] for k in SITTER_FIELDS], active, top + 1, now, who))
        sitter_id = cur.lastrowid
    else:
        con.execute(
            "UPDATE sitters SET name=?, phone=?, email=?, rate=?, how_we_know=?, kids_ok=?, "
            "availability=?, notes=?, active=?, updated_at=?, updated_by=? WHERE id=?",
            (*[clean[k] for k in SITTER_FIELDS], active, now, who, sitter_id))
    con.commit()
    return dict(con.execute("SELECT * FROM sitters WHERE id=?", (sitter_id,)).fetchone())


def move_sitter(con: sqlite3.Connection, sitter_id: int, direction: int) -> None:
    """Rotation order: who to call first. Swaps sort with the neighbor."""
    rows = roster(con)
    ids = [r["id"] for r in rows if r["active"]]
    if sitter_id not in ids:
        return
    i = ids.index(sitter_id)
    j = i + direction
    if j < 0 or j >= len(ids):
        return
    ids[i], ids[j] = ids[j], ids[i]
    for k, sid in enumerate(ids):
        con.execute("UPDATE sitters SET sort=? WHERE id=?", (k + 1, sid))
    con.commit()


def delete_sitter(con: sqlite3.Connection, sitter_id: int) -> int:
    """Removes the roster row; days that pointed at them keep coverage=Babysitter with the
    sitter cleared, so the day still reads as covered-by-a-sitter, not as unassigned."""
    con.execute("UPDATE coverage_plan SET sitter_id=NULL WHERE sitter_id=?", (sitter_id,))
    n = con.execute("DELETE FROM sitters WHERE id=?", (sitter_id,)).rowcount
    con.commit()
    return n


def booked_days(con: sqlite3.Connection, today: str) -> list[dict]:
    """Upcoming days the board has assigned to a babysitter, named or not yet named."""
    return [dict(r) for r in con.execute(
        "SELECT p.day, p.sitter_id, s.name AS sitter, p.note, p.gcal_event_id "
        "FROM coverage_plan p LEFT JOIN sitters s ON s.id = p.sitter_id "
        "WHERE p.coverage=? AND p.day >= ? ORDER BY p.day", (BABYSITTER, today))]


def plans_with_sitters(con: sqlite3.Connection) -> dict:
    """coverage_plan rows with the sitter's name and phone resolved -- what days_off.board
    reads, so the calendar description can name who is coming."""
    out = {}
    for r in con.execute(
            "SELECT p.*, s.name AS sitter, s.phone AS sitter_phone FROM coverage_plan p "
            "LEFT JOIN sitters s ON s.id = p.sitter_id"):
        out[r["day"]] = dict(r)
    return out


def workbook_bytes(con: sqlite3.Connection, today: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sheet(ws, cols, rows, widths):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C4974C")
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Sitters"
    sheet(ws, ["Order", "Name", "Phone", "Email", "Rate", "How we know them", "Kids",
               "Availability", "Notes", "Active"],
          [[i + 1, r["name"], r["phone"], r["email"], r["rate"], r["how_we_know"], r["kids_ok"],
            r["availability"], r["notes"], "yes" if r["active"] else "no"]
           for i, r in enumerate(roster(con))],
          [7, 22, 16, 26, 18, 24, 8, 26, 40, 8])
    ws2 = wb.create_sheet("Sitter sheet")
    sheet(ws2, ["Section", "Item", "Value"],
          [[r["section"], r["label"], r["value"] or ""]
           for sec in info_sections(con) for r in sec["rows"]],
          [16, 36, 60])
    ws3 = wb.create_sheet("Booked days")
    sheet(ws3, ["Day", "Sitter", "Note", "On calendar"],
          [[d["day"], d["sitter"] or "(not named yet)", d["note"] or "",
            "yes" if d["gcal_event_id"] else "no"] for d in booked_days(con, today)],
          [12, 22, 40, 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- self-test


def _self_test() -> int:
    fails = 0

    def check(name, ok):
        nonlocal fails
        print(("  PASS  " if ok else "  FAIL  ") + name)
        if not ok:
            fails += 1

    # The example family: Sam + Jordan, Ava and Leo.
    family.use_example()
    bind_household()

    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript(db.SCHEMA)

    n = seed_info(con)
    total = sum(len(l) for _, l in SECTIONS)
    check("seed writes every label once", n == total)
    check("re-seed adds nothing", seed_info(con) == 0)
    secs = info_sections(con)
    check("sections keep their order", [s["section"] for s in secs] == [s for s, _ in SECTIONS])
    ped = next(s for s in secs if s["section"] == "Pediatrician")
    check("pediatrician values start EMPTY", all(not r["value"] for r in ped["rows"]))
    pc = next(r for s in secs if s["section"] == "Emergency" for r in s["rows"] if r["label"] == "Poison Control")
    check("poison control is the one prefilled constant", "1-800-222-1222" in (pc["value"] or ""))
    row = set_info(con, ped["rows"][0]["id"], "  Oak Pediatrics ", "Sam")
    check("set_info trims and records who", row["value"] == "Oak Pediatrics" and row["updated_by"] == "Sam")
    check("re-seed keeps the entered value", seed_info(con) == 0 and
          con.execute("SELECT value FROM sitter_info WHERE id=?", (row["id"],)).fetchone()[0] == "Oak Pediatrics")
    extra = add_info_row(con, "House", "Garage code", "Jordan")
    house = next(s for s in info_sections(con) if s["section"] == "House")
    check("a parent-added fact lands last in its section", house["rows"][-1]["id"] == extra["id"])
    try:
        add_info_row(con, "House", "", "Jordan")
        check("blank label refused", False)
    except ValueError:
        check("blank label refused", True)

    try:
        save_sitter(con, {"name": "  "}, "Sam")
        check("blank sitter name refused", False)
    except ValueError:
        check("blank sitter name refused", True)
    a = save_sitter(con, {"name": "Jane", "phone": "555-0100", "kids_ok": "Both"}, "Sam")
    b = save_sitter(con, {"name": "Priya", "phone": "555-0101"}, "Jordan")
    check("two sitters, rotation order = creation order",
          [r["name"] for r in roster(con)] == ["Jane", "Priya"])
    move_sitter(con, b["id"], -1)
    check("move up swaps the order", [r["name"] for r in roster(con)] == ["Priya", "Jane"])
    move_sitter(con, b["id"], -1)
    check("move past the top is a no-op", [r["name"] for r in roster(con)] == ["Priya", "Jane"])
    b2 = save_sitter(con, {"name": "Priya", "active": "0"}, "Sam", sitter_id=b["id"])
    check("inactive sorts last but is kept", b2["active"] == 0 and roster(con)[-1]["name"] == "Priya"
          and [r["name"] for r in roster(con, include_inactive=False)] == ["Jane"])
    try:
        save_sitter(con, {"name": "X", "kids_ok": "Everyone"}, "Sam")
        check("kids_ok is validated", False)
    except ValueError:
        check("kids_ok is validated", True)

    con.execute("INSERT INTO coverage_plan (day, coverage, sitter_id) VALUES ('2026-11-06', 'Babysitter', ?)", (a["id"],))
    con.execute("INSERT INTO coverage_plan (day, coverage) VALUES ('2026-11-25', 'Babysitter')")
    con.execute("INSERT INTO coverage_plan (day, coverage) VALUES ('2026-11-26', 'Sam')")
    con.execute("INSERT INTO coverage_plan (day, coverage, sitter_id) VALUES ('2026-01-05', 'Babysitter', ?)", (a["id"],))
    con.commit()
    bd = booked_days(con, "2026-09-04")
    check("booked days: babysitter days only, upcoming only, named or not",
          [(d["day"], d["sitter"]) for d in bd] == [("2026-11-06", "Jane"), ("2026-11-25", None)])
    plans = plans_with_sitters(con)
    check("plans resolve the sitter's name and phone",
          plans["2026-11-06"]["sitter"] == "Jane" and plans["2026-11-06"]["sitter_phone"] == "555-0100"
          and plans["2026-11-26"]["sitter"] is None)
    delete_sitter(con, a["id"])
    check("deleting a sitter keeps the day as Babysitter, un-named",
          con.execute("SELECT coverage, sitter_id FROM coverage_plan WHERE day='2026-11-06'").fetchone()[:] == ("Babysitter", None))

    data = workbook_bytes(con, "2026-09-04")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    check("workbook has the three sheets with rows",
          wb.sheetnames == ["Sitters", "Sitter sheet", "Booked days"]
          and wb["Sitter sheet"].max_row == total + 2 and wb["Sitters"].max_row == 2)

    # The calendar title and description name the sitter.
    import days_off
    days_off.bind_household()
    block = {"kind": days_off.FULL, "full_kids": ["Ava"], "early_kids": [], "coverage": "Babysitter",
             "sitter": "Jane", "sitter_phone": "555-0100", "days": [], "start": "2026-11-06", "end": "2026-11-06"}
    check("event title names the sitter, not 'Babysitter off'",
          days_off.event_summary(block) == "Sitter Jane — Ava (school closed)")
    check("description carries the number", "Jane (555-0100)" in days_off.event_description(block))
    block["sitter"] = None
    check("unnamed sitter still gets a valid title", days_off.event_summary(block) == "Sitter — Ava (school closed)")
    check("Babysitter is a coverage option", "Babysitter" in days_off.COVERAGE_OPTIONS)
    d1 = {"kind": days_off.FULL, "full_kids": ["Ava"], "early_kids": [], "coverage": "Babysitter",
          "sitter_id": 1, "sitter": "Jane", "day": "2026-11-05", "titles": {}}
    d2 = {**d1, "day": "2026-11-06", "sitter_id": 2, "sitter": "Priya"}
    check("two sitters on consecutive days are two calendar events", len(days_off.collapse_runs([d1, d2])) == 2)

    print("\n" + ("ALL PASS" if not fails else f"{fails} FAILED"))
    return 1 if fails else 0


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        sys.exit(_self_test())
    print(__doc__)
